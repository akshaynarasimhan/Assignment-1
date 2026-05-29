"""
Export leads to Excel and CSV.
"""

from __future__ import annotations
import csv
import io
from typing import Optional
from backend.models import Lead


def leads_to_csv(leads: list[Lead]) -> str:
    output = io.StringIO()
    fields = ["name", "title", "company", "country", "seniority", "function",
              "linkedin_url", "email", "source", "confidence"]
    writer = csv.DictWriter(output, fieldnames=fields, extrasaction="ignore")
    writer.writeheader()
    for lead in leads:
        writer.writerow(lead.dict())
    return output.getvalue()


def leads_to_excel(leads: list[Lead]) -> bytes:
    try:
        import openpyxl
    except ImportError:
        raise RuntimeError("openpyxl is required for Excel export: pip install openpyxl")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Leads"

    headers = ["Name", "Title", "Company", "Country", "Seniority", "Function",
               "LinkedIn URL", "Email", "Source", "Confidence"]
    fields = ["name", "title", "company", "country", "seniority", "function",
              "linkedin_url", "email", "source", "confidence"]

    # Header row with bold
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = openpyxl.styles.Font(bold=True)

    for row_idx, lead in enumerate(leads, 2):
        d = lead.dict()
        for col_idx, field in enumerate(fields, 1):
            ws.cell(row=row_idx, column=col_idx, value=d.get(field, ""))

    # Auto-width columns
    for col in ws.columns:
        max_len = max((len(str(cell.value or "")) for cell in col), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 60)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


def parse_bulk_upload(file_bytes: bytes, filename: str) -> list[dict]:
    """
    Parse an uploaded Excel or CSV file containing company rows.
    Handles messy column names flexibly.
    Returns list of {company_name, domain, leadership_url}.
    """
    filename_lower = filename.lower()

    _COMPANY_ALIASES = {
        "company_name", "company name", "company", "org", "organisation",
        "organization", "firm", "business", "account", "account name",
    }
    _DOMAIN_ALIASES = {"domain", "website", "url", "web", "site"}
    _LEADERSHIP_ALIASES = {
        "leadership_url", "leadership url", "team url", "about url", "management url"
    }

    def _match_col(header: str, aliases: set) -> bool:
        return header.strip().lower() in aliases

    if filename_lower.endswith(".csv"):
        text = file_bytes.decode("utf-8-sig", errors="replace")
        reader = csv.DictReader(io.StringIO(text))
        raw_rows = list(reader)
        headers = reader.fieldnames or []
    else:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
        ws = wb.active
        rows_iter = ws.iter_rows(values_only=True)
        raw_header = next(rows_iter, None)
        if not raw_header:
            return []
        headers = [str(h or "").strip() for h in raw_header]
        raw_rows = [dict(zip(headers, (str(c or "").strip() for c in row))) for row in rows_iter]

    # Map column names
    company_col = next((h for h in headers if _match_col(h, _COMPANY_ALIASES)), None)
    domain_col = next((h for h in headers if _match_col(h, _DOMAIN_ALIASES)), None)
    leadership_col = next((h for h in headers if _match_col(h, _LEADERSHIP_ALIASES)), None)

    if not company_col:
        raise ValueError(
            f"Could not find a company name column. "
            f"Expected one of: {sorted(_COMPANY_ALIASES)}. "
            f"Found: {headers}"
        )

    results = []
    for row in raw_rows:
        company = str(row.get(company_col, "") or "").strip()
        if not company:
            continue
        results.append(
            {
                "company_name": company,
                "domain": str(row.get(domain_col, "") or "").strip() or None,
                "leadership_url": str(row.get(leadership_col, "") or "").strip() or None,
            }
        )
    return results
